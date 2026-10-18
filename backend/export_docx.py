"""
Word (.docx) versions of the results, built with python-docx:
  build_gap_report(...)   the gaps to review, in detail (can be uploaded again to practise them)
  build_full_report(...)  the whole session: scores explained, gaps, question-by-question feedback, conversation, fact-check, notes
"""
import io
import os
import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

MAROON = RGBColor(0x6D, 0x1A, 0x2B)
ROSE = RGBColor(0xB1, 0x2A, 0x47)
GREY = RGBColor(0x55, 0x4B, 0x4D)
CRITERIA = ["accuracy", "completeness", "consistency", "clarity"]
CRITERIA_HELP = {
    "accuracy": "Are the statements you made correct according to your notes?",
    "completeness": "Did you cover the important ideas, or were key parts missing?",
    "consistency": "Did you stay consistent, without contradicting yourself?",
    "clarity": "Was the explanation clear, well ordered and easy for a beginner to follow?",
}
GAP_TITLE = "Vocalis Gap Report"     # docx_reader looks for this exact title when a gap report is uploaded again


# ------------------------------------------------------------------ small helpers
def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _page_number(paragraph):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)


def _new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2
    for name, size, color in (("Title", 30, MAROON), ("Heading 1", 20, MAROON), ("Heading 2", 15, ROSE), ("Heading 3", 12.5, MAROON)):
        st = doc.styles[name]
        st.font.name = "Georgia"; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = color
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(18)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(14)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Vocalis  ·  page "); r.font.size = Pt(9); r.font.color.rgb = GREY
    _page_number(fp)
    return doc


def _inline(par, text, size=None, color=None):
    """Adds text to a paragraph, turning **bold** into bold runs."""
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text or "")):
        if not chunk:
            continue
        run = par.add_run(chunk)
        run.bold = bool(i % 2)
        if size: run.font.size = Pt(size)
        if color: run.font.color.rgb = color
    return par


def _label(doc, label, text, indent=0.0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indent)
    r = p.add_run(label + "  "); r.bold = True; r.font.color.rgb = MAROON
    _inline(p, text or "—")
    return p


def _bullets(doc, items, style="List Bullet"):
    for it in items:
        _inline(doc.add_paragraph(style=style), it)


def _markdown(doc, md):
    """A small Markdown subset (## headings, - bullets, 1. lists, **bold**) as Word paragraphs."""
    for raw in (md or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if re.match(r"^-{3,}$", line.strip()):
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            doc.add_heading(m.group(2).strip(), level=3)
        elif re.match(r"^\s*[-*•]\s+", line):
            _inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            _inline(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
        else:
            _inline(doc.add_paragraph(), line.strip())


def _meta_table(doc, rows):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    for k, v in rows:
        c = t.add_row().cells
        c[0].text = ""; c[1].text = ""
        r = c[0].paragraphs[0].add_run(k); r.bold = True; r.font.color.rgb = MAROON
        c[1].paragraphs[0].add_run(str(v))
        _shade(c[0], "F4E9E4")
        c[0].width = Inches(1.7); c[1].width = Inches(4.8)
    doc.add_paragraph()


def _score_table(doc, score):
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    t.autofit = False
    hdr = t.rows[0].cells
    for c, txt in zip(hdr, ("Criterion", "Score", "What the examiner noticed")):
        c.text = ""
        r = c.paragraphs[0].add_run(txt); r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c, "6D1A2B")
    for name in CRITERIA:
        item = (score.get("scores") or {}).get(name) or {}
        c = t.add_row().cells
        c[0].text = ""; c[1].text = ""; c[2].text = ""
        r = c[0].paragraphs[0].add_run(name.capitalize()); r.bold = True
        sub = c[0].add_paragraph(); sr = sub.add_run(CRITERIA_HELP[name]); sr.font.size = Pt(9); sr.font.color.rgb = GREY
        c[1].paragraphs[0].add_run(f"{item.get('score', '–')} / 10").bold = True
        c[2].paragraphs[0].add_run(item.get("justification", ""))
        c[0].width = Inches(2.3); c[1].width = Inches(0.9); c[2].width = Inches(3.3)
    doc.add_paragraph()


def _grade_word(overall):
    if overall is None: return ""
    return ("Excellent" if overall >= 85 else "Good" if overall >= 70 else "Fair" if overall >= 50 else "Needs more work")


def _gaps_of(result):
    """Detailed gaps if the AI produced them, else the short mistake strings from the score."""
    gaps = result.get("gaps")
    if gaps:
        return gaps
    return [{"concept_misunderstood": m, "correct_explanation": "", "recommended_review": "", "what_student_said": "",
             "why_it_matters": "", "practice_question": ""} for m in ((result.get("score") or {}).get("mistakes") or [])]


def _gap_blocks(doc, gaps, with_tick=True):
    if not gaps:
        p = doc.add_paragraph()
        _inline(p, "**No significant gaps were found.** Your explanation covered the material well.")
        return
    for i, g in enumerate(gaps, 1):
        doc.add_heading(f"Gap {i}: {g.get('concept_misunderstood', '').strip()}", level=2)
        if g.get("what_student_said"): _label(doc, "What you said:", g["what_student_said"])
        _label(doc, "What was missing or wrong:", g.get("concept_misunderstood", ""))
        if g.get("correct_explanation"): _label(doc, "The correct explanation:", g["correct_explanation"])
        if g.get("why_it_matters"): _label(doc, "Why it matters:", g["why_it_matters"])
        if g.get("recommended_review"): _label(doc, "How to fix it:", g["recommended_review"])
        if g.get("practice_question"): _label(doc, "Test yourself:", g["practice_question"])
        if with_tick:
            p = doc.add_paragraph(); r = p.add_run("☐  I have reviewed this gap"); r.font.color.rgb = GREY; r.italic = True


def _front(doc, title, subtitle, meta):
    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph(); r = p.add_run(subtitle); r.font.size = Pt(15); r.font.color.rgb = ROSE
    _meta_table(doc, meta)


def _meta(session, user, result, extra=()):
    overall = (result.get("score") or {}).get("overall")
    when = str(session["created_at"] or "")[:16]
    try:
        when = datetime.strptime(str(session["created_at"])[:19], "%Y-%m-%d %H:%M:%S").strftime("%d %B %Y, %H:%M")
    except Exception:  # noqa: BLE001
        pass
    rows = [("Student", user["name"]), ("Topic", session["topic"]), ("Session date", when)]
    if overall is not None:
        rows.append(("Overall score", f"{overall} / 100  ({_grade_word(overall)})"))
    return rows + list(extra)


# ------------------------------------------------------------------ 1. gap report
def build_gap_report(session, user, result):
    doc = _new_doc()
    score = result.get("score") or {}
    _front(doc, GAP_TITLE, "What to revise, and exactly how to fix it", _meta(session, user, result))

    doc.add_heading("How to use this report", level=1)
    doc.add_paragraph(
        "Vocalis asked you five questions about this topic and compared your explanations with your own notes. "
        "This report lists every idea that was wrong, vague or missing. For each gap you will find what you said, "
        "the correct explanation, why it matters, how to fix it and a question to test yourself. "
        "Work through the gaps in order, tick them off, then upload this file again in Vocalis to be tested on exactly these gaps.")

    doc.add_heading("Overall result", level=1)
    if score.get("summary"):
        doc.add_paragraph(score["summary"])
    if score.get("scores"):
        _score_table(doc, score)

    if result.get("strengths"):
        doc.add_heading("What you did well", level=1)
        _bullets(doc, result["strengths"])

    gaps = _gaps_of(result)
    doc.add_heading(f"Gaps to review ({len(gaps)})", level=1)
    _gap_blocks(doc, gaps)

    fc = (result.get("fact_check") or {}).get("findings", "")
    if fc and not fc.startswith("No suspected errors"):
        doc.add_heading("Suspected issues in your study material", level=1)
        doc.add_paragraph("These points were flagged using a web search. They are suspected, not certain, so check them against a trusted source.")
        _markdown(doc, fc)

    if result.get("study_plan"):
        doc.add_heading("Your study plan", level=1)
        _bullets(doc, result["study_plan"], "List Number")

    doc.add_heading("Practise these gaps again", level=1)
    doc.add_paragraph("Open Vocalis, choose Upload, and select this file. Each gap becomes a topic, so the next session tests exactly what you got wrong.")
    return _to_bytes(doc)


# ------------------------------------------------------------------ 2. full report
def build_full_report(session, user, result, messages, drawing_path=lambda m: None):
    doc = _new_doc()
    score = result.get("score") or {}
    asked = sum(1 for m in messages if m["role"] == "ai")
    answered = sum(1 for m in messages if m["role"] == "student")
    _front(doc, "Vocalis Session Report", "A complete record of your explanation session",
           _meta(session, user, result, [("Questions", f"{asked} asked, {answered} answered")]))

    doc.add_heading("1. Summary", level=1)
    overall = score.get("overall")
    if overall is not None:
        doc.add_paragraph(f"You scored {overall} out of 100 ({_grade_word(overall).lower()}). The score is the average of four criteria, each marked out of 10 and then scaled to 100.")
    if score.get("summary"):
        doc.add_paragraph(score["summary"])

    doc.add_heading("2. Scores explained", level=1)
    if score.get("scores"):
        _score_table(doc, score)

    if result.get("strengths"):
        doc.add_heading("3. What you did well", level=1)
        _bullets(doc, result["strengths"])

    gaps = _gaps_of(result)
    doc.add_heading(f"4. Gaps to review ({len(gaps)})", level=1)
    _gap_blocks(doc, gaps, with_tick=False)

    qf = result.get("question_feedback") or []
    if qf:
        doc.add_heading("5. Question-by-question feedback", level=1)
        for i, q in enumerate(qf, 1):
            doc.add_heading(f"Question {i}", level=2)
            _label(doc, "Asked:", q.get("question", ""))
            _label(doc, "Your answer:", q.get("answer_summary", ""))
            _label(doc, "Feedback:", q.get("feedback", ""))

    doc.add_heading("6. The conversation", level=1)
    doc.add_paragraph("Everything that was said, in order. Whiteboard drawings are shown as the AI understood them.")
    n = 0
    for m in messages:
        if m["role"] == "ai":
            n += 1
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
            r = p.add_run(f"Question {n}:  "); r.bold = True; r.font.color.rgb = MAROON
            _inline(p, m["text"])
        else:
            text = m["text"] or ""
            drawing = ""
            if "[Whiteboard drawing]" in text:
                text, drawing = text.split("[Whiteboard drawing]", 1)
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
            r = p.add_run("You:  "); r.bold = True; r.font.color.rgb = ROSE
            p.add_run(text.strip() or "(no spoken or typed answer)")
            path = drawing_path(m)
            if path and os.path.exists(path):
                try:
                    doc.add_picture(path, width=Inches(4.6))
                    doc.paragraphs[-1].paragraph_format.left_indent = Inches(0.3)
                except Exception:  # noqa: BLE001
                    pass
            if drawing.strip():
                q = doc.add_paragraph(); q.paragraph_format.left_indent = Inches(0.3)
                rr = q.add_run("Whiteboard:  "); rr.bold = True; rr.font.color.rgb = GREY
                q.add_run(drawing.strip()).italic = True

    fc = result.get("fact_check") or {}
    doc.add_heading("7. Fact-check", level=1)
    if fc.get("error"):
        doc.add_paragraph(f"The fact-check could not run this time ({fc['error']}).")
    elif fc.get("findings"):
        if fc["findings"].startswith("No suspected errors"):
            doc.add_paragraph("No suspected factual errors were found in your notes or in what you said.")
        else:
            doc.add_paragraph("Suspected issues found with a web search. Treat them as prompts to double-check, not as certain errors.")
            _markdown(doc, fc["findings"])
        if fc.get("sources"):
            doc.add_heading("Sources", level=3)
            for s in fc["sources"]:
                doc.add_paragraph(f"{s.get('title') or s.get('url')} — {s.get('url')}", style="List Bullet")
    else:
        doc.add_paragraph("No fact-check was produced for this session.")

    notes = result.get("notes") or {}
    doc.add_heading("8. Revision notes", level=1)
    if notes.get("notes"):
        _markdown(doc, notes["notes"])
        if notes.get("notes_translated") and notes.get("language", "English") != "English":
            doc.add_heading(f"Notes in {notes['language']}", level=2)
            _markdown(doc, notes["notes_translated"])
    else:
        doc.add_paragraph("No revision notes were produced for this session.")

    if result.get("study_plan"):
        doc.add_heading("9. Your study plan", level=1)
        _bullets(doc, result["study_plan"], "List Number")
    return _to_bytes(doc)


def _to_bytes(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
