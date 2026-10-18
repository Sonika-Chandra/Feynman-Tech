"""
Tests the whole API without a browser.   Run from the repo root:
    python backend/test_api.py            # uses the fake AI: free, instant
    python backend/test_api.py --real     # uses Gemini: costs about 12 requests of your daily quota
"""
import io
import os
import sys
import tempfile

REAL = "--real" in sys.argv
os.environ["USE_FAKE_AI"] = "0" if REAL else "1"   # overrides anything in .env
print("MODE:", "REAL Gemini (uses your quota)" if REAL else "FAKE AI (free)", "\n")
os.environ["DATABASE_PATH"] = os.path.join(tempfile.mkdtemp(), "test.db")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app import app  # noqa: E402

NOTES = (
    "Photosynthesis is the process by which plants turn light energy into chemical energy. It happens in "
    "chloroplasts, which contain chlorophyll. In the light-dependent reactions, chlorophyll absorbs sunlight and "
    "splits water, releasing oxygen and producing ATP and NADPH. In the Calvin cycle, the plant uses ATP and NADPH "
    "to turn carbon dioxide into glucose."
)
_raw = app.test_client()


class _Auth:
    """Every request automatically carries the login token once HDR is filled."""
    def _go(self, method, *a, **kw):
        kw.setdefault('headers', {}).update(HDR)
        return getattr(_raw, method)(*a, **kw)
    def get(self, *a, **kw): return self._go('get', *a, **kw)
    def post(self, *a, **kw): return self._go('post', *a, **kw)


client = _Auth()
failures = []
HDR = {}    # becomes {"Authorization": "Bearer ..."} after login


def check(name, condition, detail=""):
    print(("PASS  " if condition else "FAIL  ") + name + (f"   {detail}" if not condition else ""))
    if not condition:
        failures.append(name)


r = client.get("/health")
check("health", r.status_code == 200 and r.get_json()["status"] == "ok")
check("AI mode matches the flag", r.get_json()["fake_ai"] == (not REAL), str(r.get_json()))

# ------------------------------------------------------------------ login
r = client.post("/upload")
check("upload without login -> 401", r.status_code == 401)
r = client.post("/auth/register", json={"name": "Asha", "email": "bad", "password": "secret1"})
check("register with bad email -> 400", r.status_code == 400)
r = client.post("/auth/register", json={"name": "Asha", "email": "asha@example.com", "password": "123"})
check("register with short password -> 400", r.status_code == 400)
r = client.post("/auth/register", json={"name": "Asha", "email": "Asha@Example.com", "password": "secret1"})
check("register -> 201 with token", r.status_code == 201 and r.get_json().get("token"), r.get_data(as_text=True))
token = r.get_json()["token"]
r = client.post("/auth/register", json={"name": "Asha", "email": "asha@example.com", "password": "secret1"})
check("duplicate email -> 409", r.status_code == 409)
r = client.post("/auth/login", json={"email": "asha@example.com", "password": "wrong"})
check("wrong password -> 401", r.status_code == 401)
r = client.post("/auth/login", json={"email": "asha@example.com", "password": "secret1"})
check("login -> token", r.status_code == 200 and r.get_json().get("token"))
HDR["Authorization"] = "Bearer " + r.get_json()["token"]
r = client.get("/auth/me")
check("me -> user", r.status_code == 200 and r.get_json()["user"]["name"] == "Asha")

# a second user, to prove sessions are private
r = _raw.post("/auth/register", json={"name": "Ravi", "email": "ravi@example.com", "password": "secret2"})
OTHER = {"Authorization": "Bearer " + r.get_json()["token"]}

# ------------------------------------------------------------------ upload
r = client.post("/upload")
check("upload with no file -> 400", r.status_code == 400)
r = client.post("/upload", data={"file": (io.BytesIO(b"x"), "evil.exe")}, content_type="multipart/form-data")
check("upload wrong file type -> 400", r.status_code == 400)
r = client.post("/upload", data={"file": (io.BytesIO(NOTES.encode()), "notes.txt")}, content_type="multipart/form-data")
check("upload notes.txt -> 201", r.status_code == 201, r.get_data(as_text=True))
if r.status_code != 201:
    sys.exit("Cannot continue: upload failed (see the line above).")
upload_id = r.get_json()["upload_id"]

r = _raw.post("/session/start", json={"upload_id": upload_id}, headers=OTHER)
check("another user cannot use my upload -> 404", r.status_code == 404)
r = client.post("/session/start", json={"upload_id": 9999})
check("start with unknown upload -> 404", r.status_code == 404)

# ------------------------------------------------------------------ the AI opens, then 5 questions
r = client.post("/session/start", json={"upload_id": upload_id})
check("start session -> 201", r.status_code == 201, r.get_data(as_text=True))
if r.status_code != 201:
    sys.exit("Cannot continue: session start failed (see the line above).")
start = r.get_json()
sid = start["session_id"]
check("AI asks question 1 first", start.get("next_question") and start.get("round") == 1 and start.get("total_rounds") == 5, str(start))

r = _raw.get(f"/session/{sid}", headers=OTHER)
check("another user cannot read my session -> 404", r.status_code == 404)

r = client.post("/session/answer", json={"session_id": sid, "text": ""})
check("empty answer -> 400", r.status_code == 400)

asked, last, loop_error, saw_whiteboard, saw_diagram = 1, None, None, False, False
for i in range(8):
    if i == 1:   # answer question 2 (a diagram question in fake mode) with a drawing
        import base64
        buf = io.BytesIO(base64.b64decode(  # a 1x1 PNG, enough to test the upload path
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="))
        r = client.post("/session/answer", data={"session_id": str(sid), "text": "Here is my diagram.",
                                                  "drawing": (buf, "board.png", "image/png")},
                        content_type="multipart/form-data")
    else:
        r = client.post("/session/answer", json={"session_id": sid, "text": f"My answer number {i + 1}: water is split and oxygen is released."})
    if r.status_code != 200:
        loop_error = f"HTTP {r.status_code}: {r.get_data(as_text=True).strip()}"
        break
    last = r.get_json()
    if i == 1:
        check("drawing is described by the AI", bool(last.get("drawing_description")), str(last))
    if last["should_stop"]:
        break
    asked += 1
    saw_whiteboard = saw_whiteboard or last.get("needs_whiteboard", False)
    saw_diagram = saw_diagram or bool(last.get("diagram_svg"))
check("answer endpoint works", loop_error is None, loop_error)
check("loop stops by itself", last is not None and last["should_stop"], str(last))
check("exactly 5 questions asked", asked == 5, f"asked={asked}")
if not REAL:
    check("a whiteboard question was flagged", saw_whiteboard)
    check("a diagram to label was sent as safe SVG", saw_diagram)

r = client.post("/session/answer", json={"session_id": sid, "text": "one more"})
check("answer after finished -> 409", r.status_code == 409)

r = client.get(f"/session/{sid}")
info = r.get_json()
msgs = info["messages"]
check("messages saved in order, AI first", len(msgs) == 10 and msgs[0]["role"] == "ai" and msgs[1]["role"] == "student", f"{len(msgs)} msgs")
check("diagram is stored with the question (for resume)", any(m.get("diagram_svg", "").startswith("<svg") for m in msgs if m["role"] == "ai"))
drawn = [m for m in msgs if m["has_drawing"]]
check("drawing message flagged", len(drawn) == 1)
r = client.get(f"/message/{drawn[0]['id']}/drawing")
check("drawing can be fetched back", r.status_code == 200 and r.data[:4] == b"\x89PNG")
r = _raw.get(f"/message/{drawn[0]['id']}/drawing", headers=OTHER)
check("another user cannot fetch my drawing", r.status_code == 404)

# ------------------------------------------------------------------ results
r = client.get(f"/session/{sid}/report")
check("report before result -> 404", r.status_code == 404)

r = client.get(f"/session/{sid}/result?mode=score")
pkg = r.get_json() if r.status_code == 200 else {}
check("score result -> 200", r.status_code == 200, r.get_data(as_text=True)[:300])
check("result has score fields", set(pkg.get("score", {}).get("scores", {})) == {"accuracy", "completeness", "consistency", "clarity"})
check("result has overall 0-100", 0 <= pkg.get("score", {}).get("overall", -1) <= 100)
check("score mode skips the report", not pkg.get("report_markdown"))

r = client.get(f"/session/{sid}/result")
pkg = r.get_json() if r.status_code == 200 else {}
check("full result has the gap report", r.status_code == 200 and bool(pkg.get("report_markdown")), r.get_data(as_text=True)[:300])

r = client.get(f"/session/{sid}/report")
body = r.get_data(as_text=True)
check("report has fixed headings", all(h in body for h in ("### Concept Misunderstood", "### Correct Explanation", "### Recommended Review")))
r = client.get(f"/session/{sid}/report?download=1")
check("report download header", "attachment" in r.headers.get("Content-Disposition", ""))

r = client.get("/sessions")
hist = r.get_json()["sessions"] if r.status_code == 200 else []
check("history lists the session with its score", len(hist) == 1 and hist[0]["overall"] is not None, str(hist))
r = _raw.get("/sessions", headers=OTHER)
check("other user's history is empty", r.get_json()["sessions"] == [])

# Word downloads
import docx as _docx
r = client.get(f"/session/{sid}/export?kind=gaps")
check("gap report .docx downloads", r.status_code == 200 and "wordprocessingml" in r.headers.get("Content-Type", "") and "attachment" in r.headers.get("Content-Disposition", ""), r.get_data(as_text=True)[:200] if r.status_code != 200 else "")
gap_bytes = r.data
gd = _docx.Document(io.BytesIO(gap_bytes))
gtext = "\n".join(p.text for p in gd.paragraphs)
check("gap docx has title, gap and study plan", "Vocalis Gap Report" in gtext and "Gap 1:" in gtext and "Test yourself" in gtext and "Your study plan" in gtext, gtext[:300])
r = client.get(f"/session/{sid}/export?kind=full")
check("full report .docx downloads", r.status_code == 200 and r.data[:2] == b"PK")
fd = _docx.Document(io.BytesIO(r.data))
ftext = "\n".join(p.text for p in fd.paragraphs)
check("full docx has conversation, scores and notes", all(x in ftext for x in ("The conversation", "Question 1:", "Revision notes", "Fact-check")) and len(fd.tables) >= 2, ftext[:300])
r = client.get(f"/session/{sid}/export?kind=nope")
check("export rejects unknown kind", r.status_code == 400)
r = client.post("/upload", data={"file": (io.BytesIO(gap_bytes), "gap-report.docx")}, content_type="multipart/form-data")
check("downloaded gap report (.docx) can be uploaded again", r.status_code == 201 and any(t.startswith("Mistake 1") for t in r.get_json().get("topics", [])), r.get_data(as_text=True)[:300])

# the saved report can be uploaded again as study material
r = client.post("/upload", data={"file": (io.BytesIO(body.encode()), "mistake-report.md")}, content_type="multipart/form-data")
check("saved report can be re-uploaded", r.status_code == 201, r.get_data(as_text=True)[:200])

# ------------------------------------------------------------------ real files (Role 4) through the API
SAMPLES = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "test-material")
if os.path.isdir(SAMPLES):
    for fname, expect in [("photosynthesis.pdf", "The Calvin cycle"), ("photosynthesis.pptx", "The Calvin cycle"),
                          ("photosynthesis.docx", "The Calvin cycle"), ("photosynthesis.md", "The Calvin cycle"),
                          ("flowchart.png", "Image: flowchart")]:
        with open(os.path.join(SAMPLES, fname), "rb") as fh:
            r = client.post("/upload", data={"file": (io.BytesIO(fh.read()), fname)}, content_type="multipart/form-data")
        body = r.get_json() or {}
        check(f"upload {fname} -> topics", r.status_code == 201 and expect in body.get("topics", []), r.get_data(as_text=True)[:200])
        if fname == "photosynthesis.pdf" and r.status_code == 201:
            check("upload returns a preview for each topic", all(x["preview"] for x in body["sections"]))
            pdf_upload = body["upload_id"]

    r = client.post("/session/start", json={"upload_id": pdf_upload, "topic": "The Calvin cycle"})
    check("start session on one topic", r.status_code == 201 and r.get_json()["topic"] == "The Calvin cycle")
    topic_sid = r.get_json()["session_id"]
    import database as _db
    check("session keeps only that topic's text", "stroma" in _db.get_session(topic_sid)["material_text"] and "thylakoid" not in _db.get_session(topic_sid)["material_text"])

    r = client.post("/upload", data={"file": (io.BytesIO(b"x"), "old.doc")}, content_type="multipart/form-data")
    check("upload .doc -> friendly message", r.status_code == 400 and ".docx" in r.get_json()["error"], r.get_data(as_text=True))

    r = client.post("/session/answer", data={"session_id": str(topic_sid), "audio": (io.BytesIO(b"RIFF....WAVE"), "answer.wav", "audio/wav")}, content_type="multipart/form-data")
    check("audio answer is transcribed", r.status_code == 200 and r.get_json()["transcript"], r.get_data(as_text=True)[:200])

from svg_clean import safe_svg  # noqa: E402
evil = '<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script><rect onclick="x()" width="5" height="5"/><image href="http://evil/x.png"/></svg>'
cleaned = safe_svg(evil)
check("unsafe SVG parts are removed", "script" not in cleaned and "onclick" not in cleaned and "image" not in cleaned and "<rect" in cleaned, cleaned)
check("non-SVG text is rejected", safe_svg("hello") == "" and safe_svg(None) == "")

r = client.post("/tts", json={"text": "Hello"}, headers=HDR) if "HDR" in globals() else client.post("/tts", json={"text": "Hello"})
check("/tts in demo mode -> clean 503 JSON (browser voice takes over)", r.status_code == 503 and "error" in r.get_json())
r = client.post("/tts", json={})
check("/tts rejects empty text / no login", r.status_code in (400, 401))

r = client.get("/")
check("home page is served", r.status_code == 200 and (b"Vocalis" in r.data or b"frontend" in r.data))

r = client.post("/auth/logout")
check("logout -> ok", r.status_code == 200)
r = client.get("/auth/me")
check("token no longer works after logout", r.status_code == 401)

r = client.get("/nope")
check("unknown route -> JSON 404", r.status_code == 404 and "error" in r.get_json())

print("\n" + ("ALL TESTS PASSED" if not failures else f"{len(failures)} FAILED: {failures}"))
sys.exit(1 if failures else 0)
