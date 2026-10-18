"""PowerPoint -> one topic per slide (titled by the slide title)."""
from . import topics
from .errors import ExtractionError


def _shape_texts(shape, skip_title=None):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            out += _shape_texts(sub, skip_title)
        return out
    if skip_title is not None and shape.shape_id == skip_title.shape_id:
        return out
    if shape.has_text_frame:
        txt = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
        if txt:
            out.append(txt)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(" | ".join(cells))
    return out


def read(path):
    from pptx import Presentation
    try:
        prs = Presentation(path)
    except Exception:
        raise ExtractionError("This PowerPoint file could not be opened. Save it as .pptx and try again.")
    secs = []
    for n, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        parts = []
        for shape in slide.shapes:
            parts += _shape_texts(shape, title_shape)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Speaker notes: " + notes)
        body = "\n".join(parts).strip()
        if body or title:
            secs.append({"title": title or f"Slide {n}", "text": body or title})
    return topics.finalize(secs)
